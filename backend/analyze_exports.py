import docx
import pptx
import openpyxl
import os

base = 'd:/ECC/ECC/AI Venture Studio/exports/'

print('=== 1. WORD (.DOCX) ANALYSIS ===')
doc = docx.Document(base + 'Zone7_AI_DOCX.docx')
print('Paragraphs:', len(doc.paragraphs))
print('Tables:', len(doc.tables))
headings = [p.text for p in doc.paragraphs if p.style.name.startswith('Heading')]
print('Key Headings Count:', len(headings))
for h in headings[:8]:
    print('  -', h)

print('\n=== 2. POWERPOINT (.PPTX) ANALYSIS ===')
prs = pptx.Presentation(base + 'Zone7_AI_PPTX.pptx')
print('Total Slides:', len(prs.slides))
for i, s in enumerate(prs.slides, start=1):
    titles = [shp.text_frame.text.replace('\n', ' ') for shp in s.shapes if shp.has_text_frame and len(shp.text_frame.text) < 80]
    first_title = titles[0] if titles else 'No title'
    print(f'  Slide {i}: {first_title[:60]}')

print('\n=== 3. PDF (.PDF) ANALYSIS ===')
pdf_path = base + 'Zone7_AI_PDF.pdf'
print('File Size:', f"{os.path.getsize(pdf_path):,} bytes")
with open(pdf_path, 'rb') as f:
    header = f.read(50)
    print('Header:', header)

print('\n=== 4. EXCEL (.XLSX) ANALYSIS ===')
wb = openpyxl.load_workbook(base + 'Zone7_AI_XLSX.xlsx')
print('Sheets in Workbook:', wb.sheetnames)
for sname in wb.sheetnames:
    ws = wb[sname]
    print(f'  Sheet "{sname}": max_row={ws.max_row}, max_col={ws.max_column}')
